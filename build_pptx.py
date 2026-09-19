#!/usr/bin/env python3
"""Build the SaleGuard hackathon presentation."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
import copy

# ---------- Palette ----------
DARK_BLUE = RGBColor(0x1E, 0x3A, 0x5F)
ACCENT_BLUE = RGBColor(0x25, 0x63, 0xEB)
LIGHT_BLUE = RGBColor(0xDB, 0xEA, 0xFE)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
OFF_WHITE = RGBColor(0xF5, 0xF7, 0xFA)
DARK_TEXT = RGBColor(0x1A, 0x1A, 0x2E)
GRAY_TEXT = RGBColor(0x55, 0x60, 0x72)
GREEN = RGBColor(0x16, 0xA3, 0x4A)
RED = RGBColor(0xDC, 0x26, 0x26)
AMBER = RGBColor(0xD9, 0x77, 0x06)
LIGHT_GRAY = RGBColor(0xE7, 0xEB, 0xF0)
CARD_BG = RGBColor(0xF0, 0xF4, 0xFA)
MEDIUM_BLUE = RGBColor(0x3B, 0x5D, 0x82)

FONT = "Calibri"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def add_slide():
    return prs.slides.add_slide(BLANK)


def set_bg(slide, color):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    bg.shadow.inherit = False
    # send to back
    spTree = slide.shapes._spTree
    spTree.remove(bg._element)
    spTree.insert(2, bg._element)
    return bg


def add_rect(slide, left, top, width, height, color, line_color=None, line_width=None, shadow=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    if line_color:
        shp.line.color.rgb = line_color
        shp.line.width = line_width or Pt(1)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = shadow
    return shp


def add_rounded_rect(slide, left, top, width, height, color, line_color=None, radius=0.06):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    try:
        shp.adjustments[0] = radius
    except Exception:
        pass
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    if line_color:
        shp.line.color.rgb = line_color
        shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def add_text(slide, left, top, width, height, text, size=18, color=DARK_TEXT,
             bold=False, italic=False, align=PP_ALIGN.LEFT, font=FONT,
             anchor=MSO_ANCHOR.TOP, line_spacing=None, wrap=True):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        if line_spacing:
            p.line_spacing = line_spacing
        for run in p.runs:
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.italic = italic
            run.font.color.rgb = color
            run.font.name = font
    return tb


def add_bullets(slide, left, top, width, height, items, size=16, color=DARK_TEXT,
                 font=FONT, bullet_color=ACCENT_BLUE, space_after=10, bold_first=False,
                 anchor=MSO_ANCHOR.TOP, line_spacing=1.05):
    """items: list of (text, level) or (text, level, color_override)"""
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, item in enumerate(items):
        text, level = item[0], item[1]
        c = item[2] if len(item) > 2 else color
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        p.level = 0
        marker = "●  " if level == 0 else "‒  "
        indent = "" if level == 0 else "     "
        run = p.add_run()
        run.text = indent + marker + text
        run.font.size = Pt(size if level == 0 else size - 2)
        run.font.color.rgb = c if level == 0 else GRAY_TEXT
        run.font.name = font
        run.font.bold = (level == 0 and bold_first)
    return tb


def add_footer(slide, page_num, dark=False):
    color = RGBColor(0xB8, 0xC6, 0xDB) if dark else GRAY_TEXT
    add_text(slide, Inches(0.5), Inches(7.13), Inches(4), Inches(0.3),
              "SaleGuard", size=10, color=color, font=FONT, bold=True)
    add_text(slide, Inches(12.0), Inches(7.13), Inches(0.9), Inches(0.3),
              str(page_num), size=10, color=color, align=PP_ALIGN.RIGHT, font=FONT)


def add_accent_bar(slide, top=Inches(0), height=Inches(0.12)):
    add_rect(slide, 0, top, SW, height, ACCENT_BLUE)


def slide_header(slide, title, subtitle=None, dark_bg=False):
    """Standard header band for content slides."""
    band_h = Inches(1.15)
    add_rect(slide, 0, 0, SW, band_h, DARK_BLUE)
    add_rect(slide, 0, band_h, SW, Pt(4), ACCENT_BLUE)
    add_text(slide, Inches(0.55), Inches(0.18), Inches(11.5), Inches(0.6),
              title, size=30, color=WHITE, bold=True, font=FONT)
    if subtitle:
        add_text(slide, Inches(0.58), Inches(0.72), Inches(11.5), Inches(0.4),
                  subtitle, size=14, color=RGBColor(0xBF, 0xD3, 0xEE), font=FONT, italic=True)
    return band_h


# =========================================================
# SLIDE 1 — TITLE
# =========================================================
s = add_slide()
set_bg(s, DARK_BLUE)

# decorative accent shapes
circle = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(10.6), Inches(-1.6), Inches(4.2), Inches(4.2))
circle.fill.solid()
circle.fill.fore_color.rgb = MEDIUM_BLUE
circle.line.fill.background()
circle.shadow.inherit = False

circle2 = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(-1.8), Inches(4.8), Inches(3.6), Inches(3.6))
circle2.fill.solid()
circle2.fill.fore_color.rgb = MEDIUM_BLUE
circle2.line.fill.background()
circle2.shadow.inherit = False

# thin accent line
add_rect(s, Inches(0.9), Inches(2.55), Inches(1.4), Pt(4), ACCENT_BLUE)

add_text(s, Inches(0.9), Inches(2.75), Inches(11.5), Inches(1.3),
          "SaleGuard", size=72, color=WHITE, bold=True, font=FONT)

add_text(s, Inches(0.95), Inches(3.85), Inches(10.5), Inches(0.7),
          "AI-Powered Sales Call QA Automation", size=26, color=RGBColor(0x9C, 0xC5, 0xFB), bold=False, font=FONT)

add_text(s, Inches(0.95), Inches(4.55), Inches(10.7), Inches(0.6),
          "From manual audits to automated compliance — in seconds, not hours",
          size=17, color=RGBColor(0xD8, 0xE3, 0xF3), italic=True, font=FONT)

add_rect(s, Inches(0.9), Inches(6.75), Inches(0.55), Pt(3), ACCENT_BLUE)
add_text(s, Inches(0.9), Inches(6.85), Inches(8), Inches(0.4),
          "CIMET / econnex Hackathon 2024", size=14, color=RGBColor(0xB8, 0xC6, 0xDB), font=FONT, bold=True)


# =========================================================
# SLIDE 2 — THE PROBLEM
# =========================================================
s = add_slide()
set_bg(s, OFF_WHITE)
slide_header(s, "The Problem", "Quality assurance can't keep pace with sales volume")

# Left context box
add_rounded_rect(s, Inches(0.55), Inches(1.5), Inches(5.6), Inches(2.15), CARD_BG)
add_text(s, Inches(0.85), Inches(1.68), Inches(5.0), Inches(0.4), "CIMET's Business", size=15, color=DARK_BLUE, bold=True)
add_bullets(s, Inches(0.85), Inches(2.12), Inches(5.0), Inches(1.5), [
    ("Runs Australia's energy & broadband comparison platform", 0),
    ("Sales agents call to sell plans across 30+ retailers", 0),
    ("Thousands of sales calls placed every month", 0),
], size=14)

# Right: today's process box
add_rounded_rect(s, Inches(6.4), Inches(1.5), Inches(6.4), Inches(2.15), CARD_BG)
add_text(s, Inches(6.7), Inches(1.68), Inches(5.6), Inches(0.4), "Today's Process", size=15, color=DARK_BLUE, bold=True)
add_bullets(s, Inches(6.7), Inches(2.12), Inches(6.0), Inches(1.5), [
    ("Human auditors manually listen to 30-min recordings", 0),
    ("Findings logged into Excel checklists, retailer by retailer", 0),
    ("Each retailer enforces different compliance rules", 0),
], size=14)

# Pain points row
add_text(s, Inches(0.55), Inches(3.9), Inches(6), Inches(0.4), "Pain Points", size=17, color=DARK_TEXT, bold=True)

pains = [
    ("Doesn't Scale", "Auditors can't cover thousands of calls/month", RED),
    ("Inconsistent Scoring", "Different auditors, different standards", AMBER),
    ("Slow Turnaround", "Issues surface days after the call happened", AMBER),
    ("Compliance Risk", "Missed violations expose CIMET & retailers", RED),
]
card_w = Inches(2.95)
gap = Inches(0.2)
x = Inches(0.55)
for label, desc, color in pains:
    card = add_rounded_rect(s, x, Inches(4.4), card_w, Inches(2.1), WHITE, line_color=LIGHT_GRAY)
    add_rect(s, x, Inches(4.4), card_w, Inches(0.09), color)
    add_text(s, x + Inches(0.2), Inches(4.65), card_w - Inches(0.4), Inches(0.7),
              label, size=15, color=DARK_BLUE, bold=True)
    add_text(s, x + Inches(0.2), Inches(5.25), card_w - Inches(0.4), Inches(1.1),
              desc, size=12.5, color=GRAY_TEXT)
    x += card_w + gap

add_footer(s, 2)


# =========================================================
# SLIDE 3 — THE SOLUTION
# =========================================================
s = add_slide()
set_bg(s, OFF_WHITE)
slide_header(s, "The Solution — SaleGuard", "An end-to-end automated QA pipeline")

add_rounded_rect(s, Inches(0.55), Inches(1.45), Inches(12.25), Inches(0.95), DARK_BLUE)
add_text(s, Inches(0.9), Inches(1.45), Inches(11.6), Inches(0.95),
          "Recording in  →  AI-scored scorecard out", size=24, color=WHITE, bold=True,
          anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)

add_text(s, Inches(0.55), Inches(2.65), Inches(6), Inches(0.4),
          "Config-Driven Architecture", size=16, color=DARK_BLUE, bold=True)
add_bullets(s, Inches(0.55), Inches(3.1), Inches(5.9), Inches(1.4), [
    ("New retailer = new config rows", 0),
    ("Zero code changes required to onboard", 0),
    ("Checklists live in evaluation_config (JSONB)", 0),
], size=14.5)

add_text(s, Inches(6.9), Inches(2.65), Inches(6), Inches(0.4),
          "Gate Logic & Audit", size=16, color=DARK_BLUE, bold=True)
add_bullets(s, Inches(6.9), Inches(3.1), Inches(5.9), Inches(1.4), [
    ("Auto-submit calls that pass every check", 0),
    ("Hold failures for human review", 0),
    ("Full audit trail with override capability", 0),
], size=14.5)

# Three check types
add_text(s, Inches(0.55), Inches(4.65), Inches(6), Inches(0.4), "Three Check Types", size=17, color=DARK_TEXT, bold=True)
types = [
    ("Verbatim", "Script compliance — did the agent say the required words?", ACCENT_BLUE),
    ("Factual", "CRM data accuracy — rates, plans & dates match records", GREEN),
    ("Behavioural", "Call quality — tone, pacing, professionalism", AMBER),
]
card_w = Inches(4.0)
gap = Inches(0.15)
x = Inches(0.55)
for label, desc, color in types:
    add_rounded_rect(s, x, Inches(5.15), card_w, Inches(1.55), WHITE, line_color=LIGHT_GRAY)
    tag = add_rounded_rect(s, x + Inches(0.2), Inches(5.35), Inches(1.6), Inches(0.4), color, radius=0.5)
    add_text(s, x + Inches(0.2), Inches(5.35), Inches(1.6), Inches(0.4), label, size=13, color=WHITE,
              bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, x + Inches(0.2), Inches(5.85), card_w - Inches(0.4), Inches(0.8), desc, size=12, color=GRAY_TEXT)
    x += card_w + gap

add_footer(s, 3)


# =========================================================
# SLIDE 4 — HOW IT WORKS (PIPELINE)
# =========================================================
s = add_slide()
set_bg(s, OFF_WHITE)
slide_header(s, "How It Works", "The end-to-end pipeline, call to decision")

steps = [
    ("1", "Sales Call\nRecording", "Upload or record\nin-browser", DARK_BLUE),
    ("2", "Transcription", "Deepgram Nova-2 with\nspeaker diarization", ACCENT_BLUE),
    ("3", "AI Scoring", "LLM evaluates each\nretailer checklist item", ACCENT_BLUE),
    ("4", "Gate Decision", "auto_submit / held_critical /\nheld_low_conf / held_sample", AMBER),
]

box_w = Inches(2.55)
box_h = Inches(2.0)
gap = Inches(0.35)
start_x = Inches(0.5)
y = Inches(2.15)

centers = []
for i, (num, title, desc, color) in enumerate(steps):
    x = start_x + i * (box_w + gap)
    card = add_rounded_rect(s, x, y, box_w, box_h, WHITE, line_color=LIGHT_GRAY)
    # number badge
    badge = s.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.15), y + Inches(0.15), Inches(0.5), Inches(0.5))
    badge.fill.solid()
    badge.fill.fore_color.rgb = color
    badge.line.fill.background()
    badge.shadow.inherit = False
    tf = badge.text_frame
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = num
    r.font.size = Pt(18); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name = FONT

    add_text(s, x + Inches(0.2), y + Inches(0.8), box_w - Inches(0.4), Inches(0.6),
              title, size=15.5, color=DARK_BLUE, bold=True, align=PP_ALIGN.LEFT)
    add_text(s, x + Inches(0.2), y + Inches(1.35), box_w - Inches(0.4), Inches(0.6),
              desc, size=11, color=GRAY_TEXT, align=PP_ALIGN.LEFT)

    centers.append((x, y, box_w, box_h))
    if i < len(steps) - 1:
        arrow_x = x + box_w
        arrow = s.shapes.add_shape(MSO_SHAPE.CHEVRON, arrow_x + Inches(0.02), y + box_h/2 - Inches(0.18),
                                     gap - Inches(0.04), Inches(0.36))
        arrow.fill.solid()
        arrow.fill.fore_color.rgb = ACCENT_BLUE
        arrow.line.fill.background()
        arrow.shadow.inherit = False

# Outcome row (auto-submit / human review)
last_x, last_y, last_w, last_h = centers[-1]
arrow2_x = last_x + last_w
arrow = s.shapes.add_shape(MSO_SHAPE.CHEVRON, arrow2_x + Inches(0.02), last_y + last_h/2 - Inches(0.18),
                             Inches(0.5) - Inches(0.04), Inches(0.36))
arrow.fill.solid(); arrow.fill.fore_color.rgb = ACCENT_BLUE; arrow.line.fill.background(); arrow.shadow.inherit = False

outcome_x = last_x + last_w + Inches(0.5)
outcome_w = Inches(2.1)

# Auto-submit box (top)
add_rounded_rect(s, outcome_x, y - Inches(0.05), outcome_w, Inches(0.9), RGBColor(0xE7, 0xF7, 0xEE), line_color=GREEN)
add_text(s, outcome_x + Inches(0.1), y + Inches(0.08), outcome_w - Inches(0.2), Inches(0.3),
          "Auto-Submit", size=14, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
add_text(s, outcome_x + Inches(0.1), y + Inches(0.42), outcome_w - Inches(0.2), Inches(0.4),
          "Clean sale approved", size=10.5, color=GRAY_TEXT, align=PP_ALIGN.CENTER)

# Human review box (bottom)
add_rounded_rect(s, outcome_x, y + Inches(1.1), outcome_w, Inches(0.9), RGBColor(0xFD, 0xEE, 0xE1), line_color=AMBER)
add_text(s, outcome_x + Inches(0.1), y + Inches(1.23), outcome_w - Inches(0.2), Inches(0.3),
          "Human Review Queue", size=13, color=AMBER, bold=True, align=PP_ALIGN.CENTER)
add_text(s, outcome_x + Inches(0.1), y + Inches(1.57), outcome_w - Inches(0.2), Inches(0.4),
          "Held for manual check", size=10.5, color=GRAY_TEXT, align=PP_ALIGN.CENTER)

add_text(s, Inches(0.5), Inches(4.5), Inches(12), Inches(0.5),
          "Every recording flows through the same automated pipeline — regardless of retailer or agent.",
          size=13, color=GRAY_TEXT, italic=True, align=PP_ALIGN.CENTER)

add_footer(s, 4)


# =========================================================
# SLIDE 5 — SCORING ENGINE
# =========================================================
s = add_slide()
set_bg(s, OFF_WHITE)
slide_header(s, "Scoring Engine — The Brain", "Three evaluator types, driven by evaluation_config JSONB")

evals = [
    ("A", "Verbatim", "Compares agent speech against approved scripts, key phrases & match threshold", ACCENT_BLUE),
    ("B", "Factual", "Cross-references what the agent said vs CRM data — rates, plan names, dates", GREEN),
    ("C", "Behavioural", "Analyzes call patterns — dead air, rapport, pressure tactics", AMBER),
]
card_w = Inches(3.95)
gap = Inches(0.2)
x = Inches(0.55)
y = Inches(1.55)
for tag, label, desc, color in evals:
    add_rounded_rect(s, x, y, card_w, Inches(2.0), WHITE, line_color=LIGHT_GRAY)
    add_rect(s, x, y, card_w, Inches(0.09), color)
    badge = s.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.25), y + Inches(0.28), Inches(0.55), Inches(0.55))
    badge.fill.solid(); badge.fill.fore_color.rgb = color; badge.line.fill.background(); badge.shadow.inherit = False
    tf = badge.text_frame
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = tag
    r.font.size = Pt(20); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name = FONT
    add_text(s, x + Inches(0.95), y + Inches(0.35), card_w - Inches(1.2), Inches(0.5),
              f"Type {tag} — {label}", size=15.5, color=DARK_BLUE, bold=True)
    add_text(s, x + Inches(0.25), y + Inches(1.05), card_w - Inches(0.5), Inches(0.9),
              desc, size=12.5, color=GRAY_TEXT)
    x += card_w + gap

# Gate logic section
add_rounded_rect(s, Inches(0.55), Inches(3.85), Inches(12.25), Inches(2.7), DARK_BLUE)
add_text(s, Inches(0.9), Inches(4.05), Inches(6), Inches(0.4), "Gate Logic", size=18, color=WHITE, bold=True)

gate_rules = [
    ("Any critical check FAILS", "→  Held for review", RED),
    ("Low confidence score", "→  Held for review", AMBER),
    ("5% random sample of clean calls", "→  Held for calibration", ACCENT_BLUE),
    ("All clear", "→  Auto-submit", GREEN),
]
gy = Inches(4.65)
for cond, outcome, color in gate_rules:
    dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.95), gy + Inches(0.07), Inches(0.16), Inches(0.16))
    dot.fill.solid(); dot.fill.fore_color.rgb = color; dot.line.fill.background(); dot.shadow.inherit = False
    add_text(s, Inches(1.3), gy, Inches(6.0), Inches(0.4), cond, size=14.5, color=WHITE)
    add_text(s, Inches(8.0), gy, Inches(4.5), Inches(0.4), outcome, size=14.5, color=color, bold=True)
    gy += Inches(0.48)

add_footer(s, 5, dark=False)


# =========================================================
# SLIDE 6 — TECH STACK
# =========================================================
s = add_slide()
set_bg(s, OFF_WHITE)
slide_header(s, "Tech Stack", "Modern, pragmatic, config-first architecture")

stack = [
    ("Backend", "Python + FastAPI"),
    ("Frontend", "Next.js + TypeScript + shadcn/ui + Recharts"),
    ("Database", "PostgreSQL (JSONB for flexible config)"),
    ("Voice / Transcription", "Deepgram Nova-2 — diarization, word-level timestamps"),
    ("LLM — Primary", "Salesforce LLM Gateway → GPT-4o (free tier)"),
    ("LLM — Premium", "Anthropic Claude Sonnet (higher accuracy)"),
    ("Infrastructure", "Docker (PostgreSQL), local file storage"),
    ("Audio Capture", "MediaRecorder API (in-browser recording)"),
]

col_w = Inches(6.0)
row_h = Inches(1.15)
gap_x = Inches(0.25)
gap_y = Inches(0.15)
start_x = Inches(0.55)
start_y = Inches(1.5)

for i, (label, desc) in enumerate(stack):
    col = i % 2
    row = i // 2
    x = start_x + col * (col_w + gap_x)
    y = start_y + row * (row_h + gap_y)
    add_rounded_rect(s, x, y, col_w, row_h, WHITE, line_color=LIGHT_GRAY)
    add_rect(s, x, y, Inches(0.12), row_h, ACCENT_BLUE)
    add_text(s, x + Inches(0.35), y + Inches(0.14), col_w - Inches(0.6), Inches(0.35),
              label, size=14, color=DARK_BLUE, bold=True)
    add_text(s, x + Inches(0.35), y + Inches(0.52), col_w - Inches(0.6), Inches(0.55),
              desc, size=12.5, color=GRAY_TEXT)

add_footer(s, 6)


# =========================================================
# SLIDE 7 — AI MODELS DEEP DIVE
# =========================================================
s = add_slide()
set_bg(s, OFF_WHITE)
slide_header(s, "AI Models Deep Dive", "Purpose-built transcription + dual-provider LLM scoring")

# Transcription card
add_rounded_rect(s, Inches(0.55), Inches(1.5), Inches(12.25), Inches(1.95), CARD_BG)
add_text(s, Inches(0.85), Inches(1.65), Inches(6), Inches(0.4), "Transcription Model: Deepgram Nova-2",
          size=16, color=DARK_BLUE, bold=True)
add_bullets(s, Inches(0.85), Inches(2.1), Inches(11.6), Inches(1.3), [
    ("Speaker diarization — distinguishes Agent vs Customer", 0),
    ("Word-level timestamps for evidence linking", 0),
    ("PII redaction capability built in", 0),
    ("95%+ accuracy on phone audio", 0),
], size=14)

# LLM scoring section
add_text(s, Inches(0.55), Inches(3.65), Inches(6), Inches(0.4), "LLM Scoring Models", size=18, color=DARK_TEXT, bold=True)

card_w = Inches(6.0)
gap = Inches(0.25)
x = Inches(0.55)
y = Inches(4.15)
models = [
    ("GPT-4o", "via Salesforce LLM Gateway", "Default • Free tier • Good accuracy", ACCENT_BLUE),
    ("Claude Sonnet", "via Anthropic API", "Premium • Structured tool_use for guaranteed JSON schema output", DARK_BLUE),
]
for name, via, desc, color in models:
    add_rounded_rect(s, x, y, card_w, Inches(1.55), WHITE, line_color=LIGHT_GRAY)
    add_rect(s, x, y, card_w, Inches(0.09), color)
    add_text(s, x + Inches(0.25), y + Inches(0.22), card_w - Inches(0.5), Inches(0.4), name, size=17, color=DARK_BLUE, bold=True)
    add_text(s, x + Inches(0.25), y + Inches(0.65), card_w - Inches(0.5), Inches(0.3), via, size=12, color=GRAY_TEXT, italic=True)
    add_text(s, x + Inches(0.25), y + Inches(1.0), card_w - Inches(0.5), Inches(0.5), desc, size=12, color=GRAY_TEXT)
    x += card_w + gap

add_rounded_rect(s, Inches(0.55), Inches(5.9), Inches(12.25), Inches(1.0), DARK_BLUE)
add_text(s, Inches(0.85), Inches(6.0), Inches(11.6), Inches(0.35),
          "Dual-provider architecture — switch with one env variable", size=14.5, color=WHITE, bold=True)
add_text(s, Inches(0.85), Inches(6.35), Inches(11.6), Inches(0.4),
          "LLM_PROVIDER=gateway  or  LLM_PROVIDER=anthropic   •   Each check gets a dynamically built prompt from evaluation_config — never hardcoded",
          size=12, color=RGBColor(0xBF, 0xD3, 0xEE))

add_footer(s, 7)


# =========================================================
# SLIDE 8 — GUARDRAILS & COMPLIANCE
# =========================================================
s = add_slide()
set_bg(s, OFF_WHITE)
slide_header(s, "Guardrails & Compliance", "Built for trust, auditability, and safety")

items = [
    ("Consent Verification", "Scored as an explicit check — never assumed", GREEN),
    ("Override System", "Mandatory reason logging on every override", ACCENT_BLUE),
    ("Full Audit Trail", "Who overrode, when, and why — always traceable", ACCENT_BLUE),
    ("Version-Controlled Checks", "effective_from / effective_to dates on every rule", DARK_BLUE),
    ("Random Calibration Sample", "5% of clean calls held for human calibration", AMBER),
    ("PII Redaction", "Sensitive data masked in transcript display", GREEN),
    ("No Auto-Correction", "System only reports — never modifies lead data", RED),
]

col_w = Inches(5.95)
row_h = Inches(1.05)
gap_x = Inches(0.3)
gap_y = Inches(0.12)
start_x = Inches(0.55)
start_y = Inches(1.5)

for i, (label, desc, color) in enumerate(items):
    col = i % 2
    row = i // 2
    x = start_x + col * (col_w + gap_x)
    y = start_y + row * (row_h + gap_y)
    add_rounded_rect(s, x, y, col_w, row_h, WHITE, line_color=LIGHT_GRAY)
    dot = s.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.22), y + Inches(0.38), Inches(0.28), Inches(0.28))
    dot.fill.solid(); dot.fill.fore_color.rgb = color; dot.line.fill.background(); dot.shadow.inherit = False
    add_text(s, x + Inches(0.65), y + Inches(0.12), col_w - Inches(0.9), Inches(0.35),
              label, size=13.5, color=DARK_BLUE, bold=True)
    add_text(s, x + Inches(0.65), y + Inches(0.48), col_w - Inches(0.9), Inches(0.5),
              desc, size=11.5, color=GRAY_TEXT)

add_footer(s, 8)


# =========================================================
# SLIDE 9 — DEMO HIGHLIGHTS
# =========================================================
s = add_slide()
set_bg(s, OFF_WHITE)
slide_header(s, "Demo Highlights / Key Screens", "What you'll see in the walkthrough")

screens = [
    ("Live Analytics Dashboard", "Pass rates, critical fails & trends in real time"),
    ("Lead Detail View", "Audio player + transcript viewer + scorecard side-by-side"),
    ("Record → Transcribe → Score", "One flow — record in-browser, transcribe, and score instantly"),
    ("Admin Review Queue", "Held leads triaged; retailer & check CRUD management"),
    ("Evidence-Linked Playback", "Click evidence → audio seeks to the exact moment"),
]

y = Inches(1.55)
for i, (title, desc) in enumerate(screens):
    row_h = Inches(0.98)
    card = add_rounded_rect(s, Inches(0.55), y, Inches(12.25), row_h, WHITE, line_color=LIGHT_GRAY)
    num_box = add_rect(s, Inches(0.55), y, Inches(0.75), row_h, ACCENT_BLUE)
    add_text(s, Inches(0.55), y, Inches(0.75), row_h, str(i + 1), size=22, color=WHITE, bold=True,
              align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(1.55), y + Inches(0.13), Inches(10.9), Inches(0.4), title, size=15.5, color=DARK_BLUE, bold=True)
    add_text(s, Inches(1.55), y + Inches(0.52), Inches(10.9), Inches(0.4), desc, size=12.5, color=GRAY_TEXT)
    y += row_h + Inches(0.13)

add_footer(s, 9)


# =========================================================
# SLIDE 10 — THANK YOU
# =========================================================
s = add_slide()
set_bg(s, DARK_BLUE)

circle = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(9.8), Inches(4.6), Inches(4.6), Inches(4.6))
circle.fill.solid(); circle.fill.fore_color.rgb = MEDIUM_BLUE; circle.line.fill.background(); circle.shadow.inherit = False
circle2 = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(-2.2), Inches(-2.0), Inches(3.8), Inches(3.8))
circle2.fill.solid(); circle2.fill.fore_color.rgb = MEDIUM_BLUE; circle2.line.fill.background(); circle2.shadow.inherit = False

add_rect(s, Inches(0.9), Inches(2.35), Inches(1.4), Pt(4), ACCENT_BLUE)
add_text(s, Inches(0.9), Inches(2.55), Inches(11), Inches(1.0),
          "Thank You", size=54, color=WHITE, bold=True, font=FONT)
add_text(s, Inches(0.95), Inches(3.55), Inches(11), Inches(0.6),
          "SaleGuard — Automating Trust in Every Sale", size=22, color=RGBColor(0x9C, 0xC5, 0xFB), italic=True, font=FONT)

add_text(s, Inches(0.95), Inches(4.85), Inches(6), Inches(0.4),
          "Built by:  Shilpi Mittal", size=16, color=WHITE, font=FONT, bold=True)
add_text(s, Inches(0.95), Inches(5.3), Inches(8), Inches(0.4),
          "GitHub:  github.com/shilpi1998/SaleGuard", size=16, color=RGBColor(0xD8, 0xE3, 0xF3), font=FONT)

add_text(s, Inches(0.9), Inches(6.85), Inches(8), Inches(0.4),
          "CIMET / econnex Hackathon 2024", size=12, color=RGBColor(0xB8, 0xC6, 0xDB), font=FONT)

prs.save("/Users/shilpi.mittal/IdeaProjects/SaleGuard/SaleGuard_Presentation.pptx")
print("Saved presentation.")
